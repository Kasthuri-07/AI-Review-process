"""
Unit Tests for Internship Automation System Modules
Tests Database, Excel Manager, Task Allocator, Email Service, and AI Evaluator.
"""

import pytest
import openpyxl

import database
from config import EXCEL_PATH
from task_allocator import TaskAllocator
from email_service import EmailService
from ai_evaluator import AIEvaluator

@pytest.fixture(scope="module", autouse=True)
def setup_test_db():
    database.init_db()

def test_database_initialization():
    """Verify that all SQLite tables and audit structures are created."""
    with database.get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table'")
        tables = [r[0] for r in cursor.fetchall()]
        assert "students" in tables
        assert "task_catalog" in tables
        assert "task_allocations" in tables
        assert "review_schedules" in tables
        assert "ai_review_feedback" in tables
        assert "email_logs" in tables
        assert "audit_logs" in tables

def test_excel_workbook_sheets_and_formatting():
    """Verify that AI_Internship_Review_Management.xlsx has all 6 sheets and formulas."""
    assert EXCEL_PATH.exists(), f"Excel file {EXCEL_PATH} should exist."
    wb = openpyxl.load_workbook(str(EXCEL_PATH), data_only=False)
    sheet_names = wb.sheetnames
    expected_sheets = [
        "Student Master", "Task Allocation", "Review Schedule",
        "AI Review Feedback", "Email Log", "Dashboard"
    ]
    for s in expected_sheets:
        assert s in sheet_names, f"Expected worksheet '{s}' not found in workbook."

    # Verify Student Master row count
    ws_stu = wb["Student Master"]
    assert ws_stu.max_row >= 101, "Expected at least 100 student rows + header."

    # Verify Dashboard sheet formulas
    ws_dash = wb["Dashboard"]
    total_stu_formula = str(ws_dash["A6"].value)
    assert "COUNTA('Student Master'!" in total_stu_formula

def test_task_allocator_idempotency_and_rules():
    """Verify automatic task allocation by track and duplicate prevention."""
    allocator = TaskAllocator()
    # Test allocating to STU-095 (which is newly assigned/active)
    ok, msg, data = allocator.allocate_task_to_student(
        student_id="STU-095",
        force_resend_email=True,
        sync_to_excel=False
    )
    assert ok is True
    assert "Allocated" in msg
    assert data["task_id"].startswith("TASK-")

    # Verify duplicate email prevention without force flag
    ok_dup, msg_dup, _ = allocator.allocate_task_to_student(
        student_id="STU-095",
        force_resend_email=False,
        sync_to_excel=False
    )
    assert "Duplicate prevented" in msg_dup

def test_email_service_rendering_and_mock_capture():
    """Verify Jinja2 template rendering and outbox file capture."""
    email_service = EmailService(mode="MOCK")
    context = {
        "student_id": "TEST-001",
        "student_name": "Test Student",
        "task_id": "TASK-AIML-101",
        "task_title": "Churn Prediction Pipeline",
        "assigned_date": "2026-10-26",
        "deadline": "2026-11-05",
        "mentor_name": "Dr. Aris Thorne",
        "mentor_email": "aris.thorne@mentor-labs.org",
        "internship_role": "AI/ML Engineer",
        "task_description": "Sample description.",
        "required_deliverables": "1. Python script",
        "task_priority": "High",
        "submission_link": "https://example.com/sub"
    }

    sent, msg = email_service.send_email(
        student_id="TEST-001",
        recipient_email="test.student@example-interns.edu",
        email_type="Unit Test Allocation",
        subject="[Test] Task Allocation",
        template_name="task_allocation.html",
        context=context,
        force_resend=True
    )
    assert sent is True
    assert "Mock email generated" in msg

def test_ai_evaluator_deterministic_and_cloud_honesty():
    """Verify AI evaluator rubric scoring and cloud access honesty guardrail."""
    evaluator = AIEvaluator(provider="MOCK")

    # 1. Test unauthenticated cloud link honesty guardrail
    cloud_result = evaluator.evaluate_submission(
        student_id="STU-040",
        task_id="TASK-AIML-101",
        is_cloud_link_only=True
    )
    assert cloud_result["score"] == 0
    assert "OAuth" in cloud_result["submission_summary"]

    # 2. Test code evaluation on structured sample code
    sample_code = """
import os
from fastapi import FastAPI
app = FastAPI()

@app.get('/predict')
def predict():
    try:
        return {'status': 'ok'}
    except Exception as e:
        return {'error': str(e)}

def test_predict():
    assert predict()['status'] == 'ok'
"""
    result = evaluator.evaluate_submission(
        student_id="STU-040",
        task_id="TASK-AIML-101",
        submission_text=sample_code
    )
    assert result["score"] >= 70
    assert "FastAPI" in result["what_was_done_correctly"] or "modular" in result["what_was_done_correctly"]
